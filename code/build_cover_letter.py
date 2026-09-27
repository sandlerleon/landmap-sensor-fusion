# -*- coding: utf-8 -*-
"""Cover letter for the Environmental Monitoring and Assessment submission; numbers from
the model output.

    python build_cover_letter.py
"""
import io
import json
import os

from docx import Document
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = r"C:\Users\Leon\Downloads\LANDMAP_EMA"
R = json.load(io.open(os.path.join(HERE, "landmap_model_results.json"), encoding="utf-8"))
E = json.load(io.open(os.path.join(HERE, "landmap_econ_results.json"), encoding="utf-8"))
RB = json.load(io.open(os.path.join(HERE, "landmap_robustness_results.json"), encoding="utf-8"))
SME = json.load(io.open(os.path.join(HERE, "landmap_sensor_map_extreme_results.json"), encoding="utf-8"))
VM = json.load(io.open(os.path.join(HERE, "landmap_voi_map_results.json"), encoding="utf-8"))
REPO_URL = "https://github.com/sandlerleon/landmap-sensor-fusion"
CODE_DOI = os.environ.get("LANDMAP_CODE_DOI")
PREPRINT_DOI = os.environ.get("LANDMAP_PREPRINT_DOI")

KLAB = {"logreg": "logistic regression", "rf": "random forest", "xgb": "XGBoost"}
ARTICLE = {"logreg": "a", "rf": "a", "xgb": "an"}
SNAME = {"mag": "magnetometer", "emi": "EMI", "gpr": "GPR", "therm": "thermal", "gas": "gas"}
best = R["best_classifier"]
pooled = R["pooled_baseline"][best]
abl = R["ablation_pooled"]
nir = R["nir_augmentation_pooled"]
sv = E["sensor_value"]
boundary = RB["voi_zero_boundary"]["combined_severity_fraction"]

TITLE = ("From Classification Accuracy to Decision Value: Multimodal Sensor Fusion for "
         "Economic Prioritization of Landfill Mining")
SUBTITLE = "A Simulation-Based Value-of-Information Analysis"

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"; st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(9); st.paragraph_format.line_spacing = 1.15
for s in doc.sections:
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1.0)


def P(t, bold=False, after=9):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(t); r.bold = bold


for line in ("Leon Sandler", "Independent Researcher", "Northbrook, Illinois, USA",
             "sandler.leon@gmail.com | ORCID 0009-0007-4584-808X"):
    P(line, after=0)
P("", after=4)
P("The Editor-in-Chief", after=0)
P("Environmental Monitoring and Assessment", after=12)
P("Dear Editor,")
P("I am submitting the manuscript \u201c%s: %s\u201d for consideration as an original "
  "research article." % (TITLE, SUBTITLE))
P("Landfill-mining projects are approved or shelved on incomplete information about "
  "underground composition, and systematic assessments of European scenarios have found "
  "the majority of them unprofitable, driven largely by uncertainty that in principle could "
  "be resolved before excavation. This manuscript takes one proposed architecture \u2014 "
  "multimodal geophysical sensing (magnetometer, electromagnetic induction, ground-"
  "penetrating radar, thermal, gas) fused by machine learning and coupled to an explicit "
  "value-of-information economic layer \u2014 and evaluates it by simulation against the "
  "specific, falsifiable questions it raises, rather than the claims originally made for it.")
P("Its main findings are:", after=4)
for t in [
    "Under a spatially blocked, non-leaky validation protocol, %s %s classifier reaches pooled "
    "F1 of %.2f (metal), %.2f (organic) and %.2f (construction and demolition debris), against "
    "%.2f (plastic) and %.2f (glass); a distance-to-nearest-training-cell diagnostic pooled "
    "over %d observations shows no accuracy decay with distance, and random per-cell "
    "splitting did not detectably inflate performance relative to spatial blocking (95%% CI "
    "including zero for all three classifiers tested) \u2014 a tested, honestly reported null "
    "result, not the leakage effect I had expected to find."
    % (ARTICLE[best], KLAB[best], pooled["per_class_f1"]["metal"],
       pooled["per_class_f1"]["organic"], pooled["per_class_f1"]["cd"],
       pooled["per_class_f1"]["plastic"], pooled["per_class_f1"]["glass"],
       R["distance_decay"]["n_total"]),
    "Sensor ablation and economic value-of-information analysis disagree about which channel "
    "matters most: removing ground-penetrating radar costs the most classification "
    "performance (macro-F1 %.3f \u2192 %.3f), but only the magnetometer shows a reliably "
    "positive incremental value of information at baseline economics (positive in %.0f%% of "
    "independent surveys). The paper's central finding is that this ranking is not fixed: a "
    "grid over ordinary metal/plastic price fluctuation never changed the winner, but "
    "deliberately extreme, still economically nameable site archetypes (polymer-rich, C&D-"
    "rich, glass-rich) shifted the top sensor to GPR or EMI \u2014 sensor value is "
    "conditional on site composition and economics, not a fixed property of the sensing "
    "architecture, and classifier choice for the economic layer (random forest vs. raw or "
    "calibrated XGBoost) was tested and found not to matter. Adding a near-infrared channel "
    "raises plastic F1 from %.2f to %.2f but leaves VOI statistically unchanged."
    % (abl["all"]["macro_f1"], abl["all_minus_gpr"]["macro_f1"], sv["mag"]["p_positive"] * 100,
       nir["baseline"]["per_class_f1"]["plastic"], nir["plus_nir"]["per_class_f1"]["plastic"]),
    "Across the baseline synthetic scenarios the simulated survey has positive median VOI "
    "(median $%.0f, break-even survey cost $%.0f); a two-dimensional map over sensing-"
    "quality and economic-adversity severity shows this conclusion is considerably more "
    "sensitive to adverse secondary-material markets than to sensor degradation itself, "
    "and gives a model-specific decision-viability boundary rather than a single "
    "degradation percentage."
    % (E["voi_averaged"]["median"], E["breakeven_survey_cost_usd_averaged"]["median"]),
    "The gas-hazard classifier reaches ROC-AUC %.2f and %.0f%% recall at a safety-oriented "
    "operating point, and the manuscript closes with a field validation and falsification "
    "plan giving pre-specified acceptance criteria for every prediction."
    % (R["hazard"]["auc"]["mean"], R["hazard"]["recall"]["mean"] * 100)]:
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(4)
    p.add_run(t)
P("I believe the work suits the journal because it couples environmental geophysical "
  "sensing to a quantified, falsifiable decision framework, addresses spatial-"
  "cross-validation methodology directly and honestly (including a negative result), and "
  "closes with a pre-registered-style field validation plan rather than asserting the "
  "architecture works. All results are clearly stated as model predictions from a synthetic, "
  "physically motivated but unvalidated simulation, and the paper's central contribution is "
  "showing where and why classification accuracy and economic decision value diverge, not "
  "asserting that the sensing architecture as proposed is ready for deployment.")
P("An earlier, non-quantitative description of the architecture was submitted by me in "
  "response to an industrial open-innovation challenge focused on landfill mining. That "
  "submission was not peer reviewed or published and contained none of the models, "
  "validation protocol, or results reported here. The manuscript is original, is not under "
  "consideration elsewhere, and has not been published. I am the sole author, I have no "
  "competing interests, and the work received no external funding. The synthetic-survey "
  "generator, classification and economic models, robustness sweep and figure generators "
  "are openly available at %s%s. "
  "Generative AI (Claude, Anthropic) assisted with literature search, model implementation "
  "and drafting; I reviewed and edited all content, verified every cited reference against "
  "its source via Crossref (correcting one incorrect DOI I had used in an earlier, "
  "non-peer-reviewed version of this material), and take full responsibility for the "
  "manuscript, as stated in its Declarations."
  % (REPO_URL, ((" and archived at https://doi.org/%s" % CODE_DOI) if CODE_DOI else "")))
if PREPRINT_DOI:
    P("A preprint of the manuscript is deposited at https://doi.org/%s, as permitted by "
      "Springer's preprint policy." % PREPRINT_DOI)
P("Thank you for your consideration.", after=12)
P("Sincerely,", after=0)
P("Leon Sandler")
os.makedirs(OUT, exist_ok=True)
doc.core_properties.author = "Leon Sandler"
path = os.path.join(OUT, "LANDMAP_EMA_Cover_Letter.docx")
doc.save(path)
print("saved", path, "| words", sum(len(p.text.split()) for p in doc.paragraphs))
