# -*- coding: utf-8 -*-
"""Repair the Rubriq-edited LANDMAP manuscript (LANDMAP_EMA_Manuscript_v2.docx) without
discarding its style edits (US spelling, de-hyphenation, commas).

Rubriq's pass dropped a data-bearing clause from the abstract (median VOI and P(VOI>0)
vanished entirely), flipped a sign ($-40 -> $40, directly contradicting the sentence's own
"becomes negative" claim), broke several sentences into ungrammatical fragments, introduced
a literal typo ("arg argues") and a word-duplication glitch ("undermodelmodelled"), changed
a section heading's meaning, and left a dangling referent ("differs from that of only").
Each is restored below; every fix must hit at least once.

    python fix_rubriq_v2.py
"""
import os
import shutil

from docx import Document

D = r"C:\Users\Leon\Downloads\LANDMAP_EMA"
RAW = os.path.join(D, "_rubriq_raw", "LANDMAP_EMA_Manuscript_v2_rubriq_raw.docx")
OUT = os.path.join(D, "LANDMAP_EMA_Manuscript_v2.docx")

FIXES = [
    # abstract: Rubriq's rewrite dropped the canonical VOI number and P(VOI>0) entirely
    ("Across the baseline synthetic scenarios and a break-even survey cost of $1451, "
     "a two-dimensional map",
     "Across the baseline synthetic scenarios and stated economic parameter distributions, "
     "pooling 40 independent survey realizations, the simulated survey has positive median "
     "VOI ($924, P(VOI>0) = 0.93) and a break-even survey cost of $1451; a two-dimensional map"),
    # dangling referent ("differs from that of only" doesn't parse)
    ("so that the hazard status is determined by constructing an ideal gas sensor; the "
     "deployed sensor differs from that of only its own instrument noise and dropout, as a "
     "real hazard label derived from independent gas testing is related to a field-deployed "
     "sensor.",
     "so that the hazard status is by construction what an ideal gas sensor would detect; the "
     "deployed sensor differs from that truth only by its own instrument noise and dropout, "
     "as a real hazard label derived from independent gas testing would relate to a "
     "field-deployed sensor."),
    # heading changed from a claim about two quantities to a claim about one
    ("5.4 Sensor value: classification accuracy is not economically valuable",
     "5.4 Sensor value: classification accuracy is not economic value"),
    # literal typo
    ("baseline economics\u2014arg argues against sizing",
     "baseline economics\u2014argues against sizing"),
    # broken sentence
    ("Recall alone that this operating point's usefulness is overstated: a detector can "
     "reach 100% recall by flagging nearly everything.",
     "Recall alone overstates this operating point's usefulness: a detector can reach 100% "
     "recall by flagging nearly everything."),
    # broken sentence (dropped "and it does so")
    ("the operating point is deliberately chosen to guarantee recall for a safety "
     "application, but at a real, the quantified precision cost that the recall figure "
     "alone does not convey.",
     "the operating point is deliberately chosen to guarantee recall for a safety "
     "application, and it does so, but at a real, quantified precision cost that the recall "
     "figure alone does not convey."),
    # subject-verb agreement ("estimates pools")
    ("this paper's one canonical baseline VOI estimates pools both 40 survey realizations "
     "in total:",
     "this paper's one canonical baseline VOI estimate pools both, 40 survey realizations "
     "in total:"),
    # sign flip: $-40 (negative, matching "becomes negative" in the same sentence) -> $40
    ("the mean VOI decreases from $1012 at baseline to $985, $868, $359 and $40 across "
     "increasing combined severity, and it becomes negative",
     "the mean VOI decreases from $1012 at baseline to $985, $868, $359 and $-40 across "
     "increasing combined severity, and it becomes negative"),
    # broken sentence ("contributes to determining X a computable quantity" has no verb)
    ("This paper contributes to determining \u201chow much would it be worth knowing\u201d "
     "a computable, falsifiable quantity for a specific sensing architecture rather than a "
     "general observation about uncertainty.",
     "This paper's contribution is to make \u201chow much would it be worth to know\u201d a "
     "computable, falsifiable quantity for a specific sensing architecture, rather than a "
     "general observation about uncertainty."),
    # word-duplication glitch in a table cell
    ("the synthetic physics undermodelmodelled real discriminability",
     "the synthetic physics undermodeled real discriminability"),
]


def all_paragraphs(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                yield from c.paragraphs


def replace_in_paragraph(p, old, new):
    runs = p.runs
    text = "".join(r.text for r in runs)
    i = text.find(old)
    if i < 0:
        return False
    j = i + len(old)
    pos, first = 0, None
    for r in runs:
        a, b = pos, pos + len(r.text)
        pos = b
        if b <= i or a >= j:
            continue
        s, e = max(i, a) - a, min(j, b) - a
        if first is None:
            r.text = r.text[:s] + new + r.text[e:]
            first = r
        else:
            r.text = r.text[:s] + r.text[e:]
    return True


def main():
    if not os.path.exists(RAW):
        os.makedirs(os.path.dirname(RAW), exist_ok=True)
        shutil.copy2(OUT, RAW)
    doc = Document(RAW)
    for old, new in FIXES:
        hits = 0
        for p in all_paragraphs(doc):
            while old in p.text:
                assert replace_in_paragraph(p, old, new)
                hits += 1
        assert hits >= 1, "no hit: %r" % old[:70]
        print("%2d  %s" % (hits, old[:70]))
    doc.core_properties.author = "Leon Sandler"
    doc.save(OUT)
    print("repaired:", OUT)


if __name__ == "__main__":
    main()
